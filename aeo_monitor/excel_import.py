"""기존 조사 엑셀(이투스247_이천기숙학원 ... TOP10.xlsx)에서 질문 목록과 기준선(출처 유형 비율)을 가져온다."""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl
import yaml

# 엑셀 4번 시트의 사이트 유형 → 이 프로그램의 유형명
BASELINE_TYPE_MAP = {
    "언론/보도자료": "언론/보도자료",
    "제3자 학원정보 플랫폼": "제3자 학원정보 플랫폼",
    "기타/무관": "기타/무관",
    "자사(앱)": "자사(앱)",
    "경쟁사 공식자료": "경쟁사 공식자료",
    "공식 홈페이지(이천캠퍼스)": "공식 홈페이지(이천캠퍼스)",
    "커뮤니티(수만휘·오르비 등)": "커뮤니티(수만휘·오르비 등)",
    "네이버 블로그": "네이버 블로그",
}


def _norm_title(t: str) -> str:
    return re.sub(r"\s+", "", str(t)).rstrip("?.!")


def import_excel(path: Path, config_dir: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=False)
    names = wb.sheetnames
    questions: dict[str, dict] = {}

    def add(text: str, category: str | None, origin: str, views=0, community="", url=""):
        text = str(text).strip()
        if not text:
            return
        key = _norm_title(text)
        q = questions.setdefault(key, {"text": text, "categories": [], "origin": origin,
                                        "views": 0, "community": community, "source_url": url})
        if category and category not in q["categories"]:
            q["categories"].append(category)
        q["views"] = max(q["views"], int(views or 0))
        if origin == "representative":
            q["origin"] = "representative"

    # 2. 카테고리별 TOP10 : 실제 게시글 제목
    ws = wb[[n for n in names if "카테고리별 TOP10" in n][0]]
    cat = None
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0]:
            cat = str(row[0]).split("(")[0].strip()
        if row[5]:
            add(row[5], cat, "community_title", row[2], row[3] or "", row[6] or "")

    # 3. 카테고리별 대표 검색어
    ws = wb[[n for n in names if "검색노출" in n and "카테고리" in n][0]]
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] and row[1] and not str(row[0]).startswith("※"):
            add(row[1], str(row[0]).strip(), "representative")

    # 4. TOP30 : 질문 + 기준선 사이트 유형 합계
    ws = wb[[n for n in names if "TOP30" in n][0]]
    mix = {}
    in_detail = False
    for row in ws.iter_rows(min_row=1, values_only=True):
        if row[0] and str(row[0]).startswith("■ 질문별"):
            in_detail = True
            continue
        if not in_detail and row[0] in BASELINE_TYPE_MAP and isinstance(row[1], (int, float)):
            mix[BASELINE_TYPE_MAP[row[0]]] = int(row[1])
        if in_detail and isinstance(row[0], int) and row[4]:
            add(row[4], None, "community_title", row[1], row[2] or "")

    qlist = []
    for i, q in enumerate(sorted(questions.values(), key=lambda x: (x["origin"] != "representative", -x["views"])), 1):
        q = {"id": f"q{i:03d}", **q}
        if not q["categories"]:
            q["categories"] = ["TOP30"]
        qlist.append(q)

    header = (
        "# 엑셀에서 자동 추출한 소비자 실제 질문 목록 (python -m aeo_monitor import-excel 로 재생성)\n"
        "# - origin: community_title = 커뮤니티 게시글 제목 원문, representative = 카테고리 대표 검색어\n"
        "# - 질문을 추가/수정해도 됩니다. 측정에서 빼려면 enabled: false\n"
    )
    with open(config_dir / "questions.yaml", "w", encoding="utf-8") as f:
        f.write(header)
        yaml.safe_dump({"questions": qlist}, f, allow_unicode=True, sort_keys=False, width=200)

    baseline = {
        "label": "엑셀 기준선 (2026.9.29, Claude WebSearch 상위 9건 × TOP30 질문)",
        "source_mix": mix,
        "official_exposed": 0,
        "note": "기준선은 검색 결과 노출 기준이며, 일일 측정은 AI 답변의 출처 기준입니다.",
    }
    with open(config_dir / "baseline.json", "w", encoding="utf-8") as f:
        json.dump(baseline, f, ensure_ascii=False, indent=2)
    return {"questions": len(qlist), "baseline_types": len(mix)}
