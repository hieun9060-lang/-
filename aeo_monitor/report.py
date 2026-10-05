"""일일 리포트 생성: HTML 대시보드, 마크다운 요약, 엑셀."""
from __future__ import annotations

import html
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from .analyze import load_run
from .storage import Store

BRANDED_CHIP = '<span class="chip tg">브랜드 질문</span>'
ENGINE_LABEL = {"claude": "Claude", "chatgpt": "ChatGPT", "gemini": "Gemini", "perplexity": "Perplexity", "mock": "DEMO"}

# 차트용 출처 그룹 (범주형 색상 8개 이내)
SOURCE_GROUPS = [
    ("공식 홈페이지", ["공식 홈페이지(이천캠퍼스)"]),
    ("자사 앱·채널", ["자사(앱)", "자사(기타 채널)"]),
    ("커뮤니티", ["커뮤니티(수만휘·오르비 등)"]),
    ("블로그", ["네이버 블로그", "기타 블로그"]),
    ("언론/보도자료", ["언론/보도자료"]),
    ("제3자 학원정보", ["제3자 학원정보 플랫폼"]),
    ("경쟁사 공식", ["경쟁사 공식자료"]),
    ("기타(영상·위키·무관)", ["동영상", "위키/백과", "기타/무관"]),
]


def e(x) -> str:
    return html.escape(str(x if x is not None else ""))


def group_mix(mix: dict | None) -> list[tuple[str, int]]:
    mix = mix or {}
    known = {t for _, ts in SOURCE_GROUPS for t in ts}
    rows = []
    for name, ts in SOURCE_GROUPS:
        n = sum(mix.get(t, 0) for t in ts)
        if name.startswith("기타"):
            n += sum(v for k, v in mix.items() if k not in known)
        rows.append((name, n))
    return rows


CSS = """
:root{color-scheme:light;--bg:#f6f5f2;--surface:#fcfcfb;--border:#e4e2dc;--text:#0b0b0b;--text2:#52514e;--muted:#8a8984;
--accent:#2a78d6;--neutral:#b9b7b0;--good:#0ca30c;--critical:#d03b3b;--track:#eeede9;
--s1:#2a78d6;--s2:#eb6834;--s3:#1baf7a;--s4:#eda100;--s5:#e87ba4;--s6:#008300;--s7:#4a3aa7;--s8:#8a8984;}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){color-scheme:dark;--bg:#111110;--surface:#1a1a19;--border:#2e2e2b;
--text:#fff;--text2:#c3c2b7;--muted:#8f8e86;--accent:#3987e5;--neutral:#5a5955;--track:#262624;
--s1:#3987e5;--s2:#d95926;--s3:#199e70;--s4:#c98500;--s5:#d55181;--s6:#008300;--s7:#9085e9;--s8:#77766f;}}
:root[data-theme="dark"]{color-scheme:dark;--bg:#111110;--surface:#1a1a19;--border:#2e2e2b;--text:#fff;--text2:#c3c2b7;--muted:#8f8e86;
--accent:#3987e5;--neutral:#5a5955;--track:#262624;--s1:#3987e5;--s2:#d95926;--s3:#199e70;--s4:#c98500;--s5:#d55181;--s6:#008300;--s7:#9085e9;--s8:#77766f;}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--text);font:15px/1.6 -apple-system,BlinkMacSystemFont,"Apple SD Gothic Neo","Pretendard","Noto Sans KR",sans-serif}
main{max-width:1180px;margin:0 auto;padding:24px 16px 64px}
h1{font-size:24px;margin:0 0 4px}h2{font-size:18px;margin:36px 0 12px}h3{font-size:15px;margin:18px 0 8px;color:var(--text2)}
.sub{color:var(--text2);font-size:13px}.demo{background:#fab219;color:#000;padding:8px 12px;border-radius:8px;margin:12px 0;font-weight:600}
.card{background:var(--surface);border:1px solid var(--border);border-radius:12px;padding:16px 18px}
.grid{display:grid;gap:12px;grid-template-columns:repeat(auto-fit,minmax(190px,1fr))}
.tile .label{font-size:13px;color:var(--text2)}.tile .val{font-size:30px;font-weight:700;font-variant-numeric:tabular-nums}
.tile .note{font-size:12px;color:var(--muted)}.up{color:var(--good)}.down{color:var(--critical)}
.cols{display:grid;gap:12px;grid-template-columns:repeat(auto-fit,minmax(340px,1fr))}
.bar-row{display:grid;grid-template-columns:minmax(110px,190px) 1fr 112px;gap:10px;align-items:center;margin:6px 0;font-size:13px}
.bar-row .name{overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
.track{background:var(--track);height:12px;border-radius:0 4px 4px 0;position:relative}
.fill{height:12px;background:var(--neutral);border-radius:0 4px 4px 0}.fill.t{background:var(--accent)}
.num{text-align:right;font-variant-numeric:tabular-nums;color:var(--text2)}
.stack{display:flex;height:18px;gap:2px;border-radius:0 4px 4px 0;overflow:hidden;background:var(--track)}
.stack span{display:block;height:100%}
.legend{display:flex;flex-wrap:wrap;gap:6px 14px;font-size:12px;color:var(--text2);margin:8px 0}
.legend i{display:inline-block;width:10px;height:10px;border-radius:2px;margin-right:5px;vertical-align:-1px}
table{border-collapse:collapse;width:100%;font-size:13px}th,td{border-bottom:1px solid var(--border);padding:7px 8px;text-align:left;vertical-align:top}
th{color:var(--text2);font-weight:600;background:var(--surface);position:sticky;top:0}
.tbl{overflow-x:auto;max-width:100%}
.ok{color:var(--good);font-weight:700}.no{color:var(--critical);font-weight:700}.half{color:#b07800;font-weight:700}
.chip{display:inline-block;font-size:11px;padding:1px 7px;border:1px solid var(--border);border-radius:999px;margin:1px 2px;color:var(--text2);white-space:nowrap}
.chip.tg{border-color:var(--accent);color:var(--accent)}
.action{border-left:4px solid var(--accent);margin:10px 0}.action.p-높음{border-left-color:var(--critical)}.action.p-중간{border-left-color:#eda100}
.action h4{margin:0 0 4px;font-size:15px}.action ul{margin:6px 0 0;padding-left:18px}.action .ev{font-size:12px;color:var(--muted)}
.brief li{margin:4px 0}.brief{white-space:normal}
details summary{cursor:pointer;color:var(--text2)}
svg text{fill:var(--text2);font-size:11px}
.toggle{float:right;font-size:12px;background:var(--surface);color:var(--text2);border:1px solid var(--border);border-radius:6px;padding:4px 8px;cursor:pointer}
@media (max-width:600px){.bar-row{grid-template-columns:92px 1fr 96px}.tile .val{font-size:24px}}
"""

THEME_JS = """
<script>
(function(){try{var t=localStorage.getItem('aeo-theme');if(t)document.documentElement.setAttribute('data-theme',t);}catch(e){}
window.toggleTheme=function(){var r=document.documentElement;var cur=r.getAttribute('data-theme')||
(matchMedia('(prefers-color-scheme: dark)').matches?'dark':'light');var n=cur==='dark'?'light':'dark';
r.setAttribute('data-theme',n);try{localStorage.setItem('aeo-theme',n);}catch(e){}};})();
</script>
"""


def _bar_rows(rows: list[tuple[str, float, str, bool]], max_val: float = 100.0) -> str:
    out = []
    for name, val, label, is_target in rows:
        w = 0 if not max_val else max(0.0, min(100.0, 100.0 * val / max_val))
        out.append(
            f'<div class="bar-row" title="{e(name)}: {e(label)}"><div class="name">{e(name)}</div>'
            f'<div class="track"><div class="fill{" t" if is_target else ""}" style="width:{w:.1f}%"></div></div>'
            f'<div class="num">{e(label)}</div></div>')
    return "".join(out)


def _stack(mix: list[tuple[str, int]]) -> str:
    total = sum(n for _, n in mix) or 1
    segs = []
    for i, (name, n) in enumerate(mix, 1):
        if n:
            segs.append(f'<span style="width:{100 * n / total:.2f}%;background:var(--s{i})" '
                        f'title="{e(name)}: {n}건 ({100 * n / total:.1f}%)"></span>')
    return f'<div class="stack">{"".join(segs)}</div>'


def _legend() -> str:
    return '<div class="legend">' + "".join(
        f'<span><i style="background:var(--s{i})"></i>{e(n)}</span>' for i, (n, _) in enumerate(SOURCE_GROUPS, 1)) + "</div>"


def _sparkline(trend: list[dict]) -> str:
    if len(trend) < 2:
        return '<p class="sub">추이는 2일 이상 측정이 쌓이면 표시됩니다.</p>'
    w, h, pad = 640, 150, 28
    xs = [pad + i * (w - 2 * pad) / (len(trend) - 1) for i in range(len(trend))]
    ys = [h - pad - (t["rate"] / 100) * (h - 2 * pad) for t in trend]
    pts = " ".join(f"{x:.1f},{y:.1f}" for x, y in zip(xs, ys))
    dots = "".join(
        f'<g><circle cx="{x:.1f}" cy="{y:.1f}" r="4" fill="var(--accent)" stroke="var(--surface)" stroke-width="2"/>'
        f'<circle cx="{x:.1f}" cy="{y:.1f}" r="12" fill="transparent"><title>{t["date"]}: {t["rate"]}% '
        + " / ".join(f"{ENGINE_LABEL.get(k, k)} {v}%" for k, v in sorted(t["engines"].items()))
        + "</title></circle></g>" for x, y, t in zip(xs, ys, trend))
    grid = "".join(
        f'<line x1="{pad}" x2="{w - pad}" y1="{h - pad - g / 100 * (h - 2 * pad):.1f}" y2="{h - pad - g / 100 * (h - 2 * pad):.1f}" '
        f'stroke="var(--border)" stroke-width="1"/><text x="2" y="{h - pad - g / 100 * (h - 2 * pad) + 4:.1f}">{g}%</text>'
        for g in (0, 50, 100))
    labels = f'<text x="{xs[0]:.1f}" y="{h - 6}" text-anchor="start">{trend[0]["date"][5:]}</text>' \
             f'<text x="{xs[-1]:.1f}" y="{h - 6}" text-anchor="end">{trend[-1]["date"][5:]}</text>'
    return (f'<svg viewBox="0 0 {w} {h}" width="100%" style="max-width:820px;display:block" role="img" aria-label="이천캠퍼스 언급률 추이">{grid}'
            f'<polyline points="{pts}" fill="none" stroke="var(--accent)" stroke-width="2"/>{dots}{labels}</svg>')


def _md_to_html(md: str) -> str:
    """AI 브리핑(간단한 마크다운) → HTML."""
    import re
    lines, out, in_list = md.splitlines(), [], False
    for ln in lines:
        s = ln.strip()
        if not s:
            continue
        body = e(re.sub(r"^[-*•]\s+|^\d+\.\s+", "", s))
        body = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", body)
        if re.match(r"^([-*•]|\d+\.)\s+", s):
            if not in_list:
                out.append("<ul class='brief'>")
                in_list = True
            out.append(f"<li>{body}</li>")
        else:
            if in_list:
                out.append("</ul>")
                in_list = False
            out.append("<p>" + re.sub(r"^#+\s*", "", body) + "</p>")
    if in_list:
        out.append("</ul>")
    return "".join(out)


def render_html(a: dict, briefing_rule: list[str], briefing_ai: str) -> str:
    import re
    k = a["kpi"]
    run = a["run"]
    engines = a["engines"]
    delta_html = ""
    if a["delta"]:
        d = a["delta"]["diff"]
        cls = "up" if d > 0 else ("down" if d < 0 else "")
        delta_html = f'<div class="note"><span class="{cls}">{"▲" if d > 0 else "▼" if d < 0 else "–"} {abs(d)}%p</span> vs {e(a["delta"]["date"])}</div>'

    tiles = f"""
<div class="grid">
 <div class="card tile"><div class="label">이천캠퍼스 언급률</div><div class="val">{k['target_rate']}%</div>{delta_html}
  <div class="note">{k['target_mentions']} / {k['core_responses']} 답변</div></div>
 <div class="card tile"><div class="label">'이투스247'만 언급 (캠퍼스 불명)</div><div class="val">{k['brand_only_rate']}%</div>
  <div class="note">{k['brand_only']}건 · 타 캠퍼스 혼동 신호 {k['campus_confusion']}건</div></div>
 <div class="card tile"><div class="label">공식 홈페이지 출처 인용률</div><div class="val">{k['official_cited_rate']}%</div>
  <div class="note">기준선(9/29) 0%</div></div>
 <div class="card tile"><div class="label">언급 시 평균 순위</div><div class="val">{k['avg_rank'] if k['avg_rank'] is not None else '–'}</div>
  <div class="note">1 = 답변에서 가장 먼저 언급</div></div>
 <div class="card tile"><div class="label">언급 시 감성</div><div class="val" style="font-size:18px;line-height:1.9">
  {' · '.join(f'{ {"positive":"긍정","negative":"부정","mixed":"혼재","neutral":"중립"}.get(s, s)} {n}' for s, n in k['sentiment'].items()) or '–'}</div>
  <div class="note">언급 주변 문맥 키워드 기준(참고용)</div></div>
</div>"""

    bold = lambda t: re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", e(t))  # noqa: E731
    brief = "<ul class='brief'>" + "".join(f"<li>{bold(x)}</li>" for x in briefing_rule) + "</ul>"
    if briefing_ai:
        brief += f"<h3>AI 해설 (Claude)</h3>{_md_to_html(briefing_ai)}"

    def rate_rows(rows, label_fn=lambda r: r["key"]):
        return _bar_rows([(label_fn(r), r["rate"], f'{r["rate"]}% ({r["hit"]}/{r["n"]})', True) for r in rows])

    sov_max = max([s["rate"] for s in a["sov"]] + [1])
    sov_rows = _bar_rows([(s["name"], s["rate"], f'{s["rate"]}%', s["is_target"]) for s in a["sov"] if s["mentions"] or s["is_target"]], sov_max)
    sov_table = "".join(
        f"<tr><td>{'<b>' if s['is_target'] else ''}{e(s['name'])}{'</b>' if s['is_target'] else ''}</td><td>{s['mentions']}</td>"
        f"<td>{s['rate']}%</td><td>{s['share']}%</td><td>{s['first']}</td><td>{s['avg_rank'] or '–'}</td></tr>" for s in a["sov"])

    tpl_rows = "".join(
        f"<tr><td>{'<b>' if t['is_target'] else ''}{e(t['name'])}{'</b>' if t['is_target'] else ''}</td><td>{t['n']}</td>"
        f"<td>{t['self_rate']}%</td><td>{t['official_rate']}%</td><td>{t['avg_cites']}</td>"
        f"<td>{', '.join(f'{e(n)} {c}' for n, c in t['top_types'])}</td>"
        f"<td>{', '.join(f'{e(n)} {c}' for n, c in t['cross_mentions']) or '–'}</td></tr>" for t in a["template_compare"])
    tpl_html = (f"""<div class="card tbl"><table><thead><tr><th>학원</th><th>답변 수</th><th>해당 학원 언급률</th><th>자사 사이트 인용률</th>
<th>평균 출처 수</th><th>주요 출처 유형</th><th>함께 언급된 학원</th></tr></thead><tbody>{tpl_rows}</tbody></table></div>
<p class="sub">같은 질문 템플릿(후기·수업·면학·시설·급식·비용·추천)에 학원명만 바꿔 질문한 결과. '자사 사이트 인용률'이 높을수록 AI가 그 학원 공식 정보를 근거로 답합니다
(brands.yaml 의 domains 정확도에 좌우됨).</p>""" if tpl_rows else '<p class="sub">학원별 동일 질문(question_templates.yaml)이 비활성화되어 있습니다.</p>')

    src = a["sources"]
    stack_rows = [("전체 출처", src["mix_all"]), ("답변 각주로 인용된 출처", src["mix_cited"]),
                  ("이천캠퍼스 언급된 답변", src["mix_when_mentioned"]), ("이천캠퍼스 미언급 답변", src["mix_when_not"])]
    if src.get("baseline"):
        stack_rows.append(("엑셀 기준선 9/29", src["baseline"]))
    stacks = "".join(
        f'<div class="bar-row"><div class="name">{e(n)}</div>{_stack(group_mix(m))}<div class="num">{sum((m or {}).values())}건</div></div>'
        for n, m in stack_rows)
    all_types = sorted({t for _, m in stack_rows for t in (m or {})})
    mix_table = "<tr><th>사이트 유형</th>" + "".join(f"<th>{e(n)}</th>" for n, _ in stack_rows) + "</tr>" + "".join(
        "<tr><td>" + e(t) + "</td>" + "".join(f"<td>{(m or {}).get(t, 0)}</td>" for _, m in stack_rows) + "</tr>" for t in all_types)

    def dom_table(rows):
        if not rows:
            return '<p class="sub">해당 없음</p>'
        return "<table><tr><th>도메인</th><th>유형</th><th>건수</th></tr>" + "".join(
            f"<tr><td>{e(d)}</td><td>{e(t)}</td><td>{n}</td></tr>" for d, n, t in rows) + "</table>"

    reasons_m = "".join(f"<tr><td>{e(t)}</td><td>{n}</td></tr>" for t, n, _ in a["reasons_mentioned"]) or "<tr><td colspan=2>–</td></tr>"
    reasons_n = "".join(f"<tr><td>{e(t)}</td><td>{n}</td></tr>" for t, n, _ in a["reasons_not"]) or "<tr><td colspan=2>–</td></tr>"

    actions = "".join(
        f"""<div class="card action p-{e(x['priority'])}"><h4>[{e(x['priority'])}] {e(x['title'])} <span class="chip">{e(x['area'])}</span></h4>
<div class="ev">근거: {e(x['reason'])} — {x['count']}건 ({x['share']}% of 답변)</div>
<ul>{''.join(f'<li>{e(d)}</li>' for d in x['detail'])}</ul>
<div class="ev">해당 질문 예: {' / '.join(e(q) for q in x['examples'])}</div></div>""" for x in a["actions"])

    from .analyze import REASONS
    qrows = []
    for q in a["questions"]:
        cells = []
        for eng in engines:
            r = q["engines"].get(eng)
            if not r:
                cells.append("<td>–</td>")
                continue
            if r["error"] and not r["mentioned"] and not r["sources"]:
                mark = f'<span class="sub" title="{e(r["error"])}">오류</span>'
            elif r["mentioned"]:
                mark = f'<span class="ok">● {r["rank"]}위</span>'
            elif r["brand_only"]:
                mark = '<span class="half">◐ 브랜드만</span>'
            else:
                mark = '<span class="no">○ 미언급</span>'
            comp = "".join(f'<span class="chip">{e(c)}</span>' for c in r["competitors"][:3])
            srcs = "".join(
                f'<div><a href="{e(u)}" target="_blank" rel="noopener">{e(d)}</a> <span class="chip{" tg" if t_ else ""}">{e(t)}</span></div>'
                for d, t, u, t_ in r["sources"])
            reasons = "".join(f"<li>{e(REASONS.get(c, c))}</li>" for c in r["reasons"])
            snippet = f'<div class="sub">“{e(r["snippet"][:140])}”</div>' if r["snippet"] else ""
            cells.append(f"<td>{mark}<div>{comp}</div><details><summary>왜? · 출처 {len(r['sources'])}</summary>"
                         f"<ul>{reasons}</ul>{snippet}{srcs}</details></td>")
        cats = "".join(f'<span class="chip">{e(c)}</span>' for c in q["categories"])
        qrows.append(f"<tr><td>{e(q['question'])}<div>{cats}{BRANDED_CHIP if q['branded'] else ''}</div></td>{''.join(cells)}</tr>")
    q_head = "".join(f"<th>{e(ENGINE_LABEL.get(x, x))}</th>" for x in engines)

    demo = '<div class="demo">DEMO — 모의 엔진으로 생성한 예시입니다. 실제 AI 측정값이 아닙니다.</div>' if run["demo"] else ""
    eng_list = ", ".join(json.loads(run["engines"]))

    return f"""<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>AI 언급 모니터 {e(run['run_date'])}</title><style>{CSS}</style>{THEME_JS}</head><body><main>
<button class="toggle" onclick="toggleTheme()">라이트/다크</button>
<h1>{e(a['target'])} · AI 챗봇 언급 리포트</h1>
<div class="sub">측정일 {e(run['run_date'])} (KST) · 엔진 {e(eng_list)} · 응답 {k['responses']}건 (오류 {k['errors']})</div>
{demo}
<h2>오늘의 핵심</h2><div class="card">{brief}</div>
<h2>핵심 지표</h2>{tiles}
<h2>언급률 추이 (최근 {len(a['trend'])}회)</h2><div class="card">{_sparkline(a['trend'])}</div>
<h2>어떤 질문에서 언급되나</h2>
<div class="cols">
 <div class="card"><h3>AI 엔진별</h3>{rate_rows(a['by_engine'], lambda r: ENGINE_LABEL.get(r['key'], r['key']))}</div>
 <div class="card"><h3>질문 유형별</h3>{rate_rows(a['by_branded'])}<h3>질문 카테고리별 (엑셀 분류)</h3>{rate_rows(a['by_category'])}</div>
</div>
<h2>경쟁 기숙학원 대비 언급 점유율</h2>
<div class="cols"><div class="card"><h3>동일한 실제 소비자 질문에서 학원별 언급률</h3>{sov_rows}</div>
<div class="card tbl"><table><thead><tr><th>학원</th><th>언급</th><th>언급률</th><th>점유율</th><th>1순위 언급</th><th>평균 순위</th></tr></thead>
<tbody>{sov_table}</tbody></table></div></div>
<h2>같은 질문, 학원별 비교</h2>{tpl_html}
<h2>출처 분석 — AI는 어디를 보고 답하나</h2>
<div class="card">{_legend()}{stacks}<details><summary>표로 보기</summary><div class="tbl"><table>{mix_table}</table></div></details>
<p class="sub">기준선은 엑셀(9/29) 웹검색 상위 노출 집계로, 측정 방식이 달라 직접 비교는 참고용입니다.</p></div>
<div class="cols" style="margin-top:12px">
 <div class="card tbl"><h3>가장 많이 인용된 도메인</h3>{dom_table(src['top_domains'])}</div>
 <div class="card tbl"><h3>이천캠퍼스를 다룬 출처 (언급의 근거)</h3>{dom_table(src['target_source_domains'])}</div>
 <div class="card tbl"><h3>경쟁사만 언급될 때의 출처</h3>{dom_table(src['domains_when_competitor'])}</div>
</div>
<h2>왜 언급됐고, 왜 언급되지 않았나</h2>
<div class="cols"><div class="card tbl"><h3>언급된 답변의 원인</h3><table>{reasons_m}</table></div>
<div class="card tbl"><h3>미언급 답변의 원인</h3><table>{reasons_n}</table></div></div>
<h2>AEO/GEO 개선과제 (우선순위)</h2>{actions or '<p class="sub">해당 없음</p>'}
<h2>질문별 상세</h2>
<p class="sub">● 이천캠퍼스 언급(답변 내 순위) · ◐ '이투스247'만 언급 · ○ 미언급. 칩은 함께 언급된 경쟁 학원, 파란 칩은 이천캠퍼스를 다룬 출처. 미언급 질문이 위에 옵니다.</p>
<div class="card tbl"><table><thead><tr><th style="min-width:240px">질문</th>{q_head}</tr></thead><tbody>{''.join(qrows)}</tbody></table></div>
<p class="sub" style="margin-top:24px">생성: aeo_monitor · 언급 판정은 brands.yaml 별칭 기준 자동 판정이며, AI 답변은 실행마다 달라질 수 있습니다.</p>
</main></body></html>"""


def render_markdown(a: dict, briefing_rule: list[str], briefing_ai: str, report_url: str = "") -> str:
    k = a["kpi"]
    run = a["run"]
    lines = [f"## {a['target']} AI 언급 리포트 — {run['run_date']}", ""]
    if run["demo"]:
        lines += ["> ⚠️ DEMO (모의 데이터)", ""]
    lines += [f"- {x}" for x in briefing_rule]
    if briefing_ai:
        lines += ["", "### AI 해설", briefing_ai]
    lines += ["", "### 엔진별 언급률", "| 엔진 | 언급률 | 언급/답변 |", "|---|---|---|"]
    lines += [f"| {ENGINE_LABEL.get(r['key'], r['key'])} | {r['rate']}% | {r['hit']}/{r['n']} |" for r in a["by_engine"]]
    lines += ["", "### 경쟁 학원 언급률 (상위 8)", "| 학원 | 언급률 | 1순위 |", "|---|---|---|"]
    lines += [f"| {'**' + s['name'] + '**' if s['is_target'] else s['name']} | {s['rate']}% | {s['first']} |" for s in a["sov"][:8]]
    lines += ["", "### 우선 개선과제"]
    lines += [f"{i}. **{x['title']}** — {x['reason']} ({x['count']}건)" for i, x in enumerate(a["actions"][:4], 1)]
    unmentioned = [q["question"] for q in a["questions"] if not any(e_["mentioned"] for e_ in q["engines"].values())]
    lines += ["", f"### 어느 엔진에서도 언급되지 않은 질문 ({len(unmentioned)}건, 일부)"]
    lines += [f"- {q}" for q in unmentioned[:10]]
    if report_url:
        lines += ["", f"📊 전체 리포트: {report_url}"]
    return "\n".join(lines) + "\n"


def write_xlsx(store: Store, a: dict, path: Path) -> None:
    from .analyze import REASONS
    from .config import load_brands
    brands = load_brands()
    wb = openpyxl.Workbook()
    head_fill = PatternFill("solid", fgColor="DDE8F7")

    def sheet(title, header, rows, widths=None):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
        for r in rows:
            ws.append(list(r))
        for i, w in enumerate(widths or [], 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        return ws

    wb.remove(wb.active)
    k = a["kpi"]
    sheet("요약", ["지표", "값"], [
        ("측정일", a["run"]["run_date"]), ("엔진", ", ".join(json.loads(a["run"]["engines"]))),
        ("DEMO 여부", "예" if a["run"]["demo"] else "아니오"),
        ("이천캠퍼스 언급률(%)", k["target_rate"]), ("언급/답변", f"{k['target_mentions']}/{k['core_responses']}"),
        ("전일 대비(%p)", a["delta"]["diff"] if a["delta"] else ""),
        ("'이투스247'만 언급 비율(%)", k["brand_only_rate"]), ("타 캠퍼스 혼동 신호(건)", k["campus_confusion"]),
        ("공식 홈페이지 인용률(%)", k["official_cited_rate"]), ("언급 시 평균 순위", k["avg_rank"] or ""),
    ], [32, 60])
    sheet("엔진·카테고리별", ["구분", "항목", "언급", "답변", "언급률(%)"],
          [("엔진", ENGINE_LABEL.get(r["key"], r["key"]), r["hit"], r["n"], r["rate"]) for r in a["by_engine"]]
          + [("질문유형", r["key"], r["hit"], r["n"], r["rate"]) for r in a["by_branded"]]
          + [("카테고리", r["key"], r["hit"], r["n"], r["rate"]) for r in a["by_category"]], [12, 28, 8, 8, 10])
    sheet("경쟁사 점유율", ["학원", "언급 수", "언급률(%)", "점유율(%)", "1순위 언급", "평균 순위"],
          [(s["name"], s["mentions"], s["rate"], s["share"], s["first"], s["avg_rank"] or "") for s in a["sov"]], [30, 10, 10, 10, 10, 10])
    sheet("학원별 동일질문", ["학원", "답변 수", "해당 학원 언급률(%)", "자사 사이트 인용률(%)", "평균 출처 수", "주요 출처 유형", "함께 언급된 학원"],
          [(t["name"], t["n"], t["self_rate"], t["official_rate"], t["avg_cites"],
            ", ".join(f"{n} {c}" for n, c in t["top_types"]), ", ".join(f"{n} {c}" for n, c in t["cross_mentions"]))
           for t in a["template_compare"]], [30, 8, 14, 14, 10, 40, 40])
    sheet("개선과제", ["우선순위", "과제", "영역", "근거(원인)", "건수", "실행 항목", "해당 질문 예"],
          [(x["priority"], x["title"], x["area"], x["reason"], x["count"], "\n".join(x["detail"]), "\n".join(x["examples"]))
           for x in a["actions"]], [8, 36, 14, 40, 8, 90, 50])

    resps = load_run(store, a["run"]["id"])
    tid = brands.target_id
    rows, crow = [], []
    for r in resps:
        m = r["mentions"].get(tid, {})
        comps = [brands.by_id(b).name for b, x in r["mentions"].items() if x["mentioned"] and b != tid]
        from .analyze import diagnose
        rows.append((r["question"], ", ".join(r["categories"]), r["origin"], "Y" if r["branded"] else "N",
                     r.get("brand_scope") or "", ENGINE_LABEL.get(r["engine"], r["engine"]), r["model"],
                     "언급" if m.get("mentioned") else ("브랜드만" if r["group_only"] else "미언급"),
                     m.get("rank") or "", m.get("sentiment") or "", ", ".join(comps),
                     len(r["citations"]), sum(1 for c in r["citations"] if c["mentions_target"]),
                     " / ".join(REASONS.get(c, c) for c in diagnose(r, brands)), r["error"] or "",
                     (r["answer"] or "")[:3000]))
        for c in r["citations"]:
            crow.append((r["question"], ENGINE_LABEL.get(r["engine"], r["engine"]), c["domain"], c["source_type"],
                         "Y" if c["cited_in_answer"] else "N", "Y" if c["mentions_target"] else "N", c["title"], c["url"]))
    ws = sheet("질문별 결과", ["질문", "카테고리", "질문 출처", "브랜드질문", "템플릿 대상", "엔진", "모델", "이천캠퍼스", "순위", "감성",
                           "함께 언급된 학원", "출처 수", "이천 다룬 출처 수", "원인", "오류", "답변"], rows,
               [40, 14, 14, 8, 14, 10, 16, 10, 6, 8, 30, 8, 10, 60, 16, 80])
    for row in ws.iter_rows(min_row=2):
        row[-1].alignment = Alignment(wrap_text=False)
    sheet("출처", ["질문", "엔진", "도메인", "사이트 유형", "답변 각주 인용", "이천캠퍼스 다룸", "제목", "URL"], crow,
          [40, 10, 24, 22, 10, 10, 50, 60])
    wb.save(path)


def write_reports(store: Store, a: dict, briefing_rule: list[str], briefing_ai: str, docs_dir: Path,
                  site_url: str = "") -> dict:
    run_date = a["run"]["run_date"]
    rep_dir = docs_dir / "reports"
    rep_dir.mkdir(parents=True, exist_ok=True)
    prefix = "demo-" if a["run"]["demo"] else ""
    html_path = rep_dir / f"{prefix}{run_date}.html"
    url = f"{site_url.rstrip('/')}/reports/{html_path.name}" if site_url else ""
    html_doc = render_html(a, briefing_rule, briefing_ai)
    html_path.write_text(html_doc, encoding="utf-8")
    md = render_markdown(a, briefing_rule, briefing_ai, url)
    md_path = rep_dir / f"{prefix}{run_date}.md"
    md_path.write_text(md, encoding="utf-8")
    xlsx_path = rep_dir / f"{prefix}{run_date}.xlsx"
    write_xlsx(store, a, xlsx_path)
    if not a["run"]["demo"]:
        (docs_dir / "index.html").write_text(html_doc, encoding="utf-8")
        (rep_dir / "latest.md").write_text(md, encoding="utf-8")
        _write_archive(rep_dir)
    return {"html": html_path, "md": md_path, "xlsx": xlsx_path, "url": url}


def _write_archive(rep_dir: Path) -> None:
    items = sorted((p.stem for p in rep_dir.glob("20*.html")), reverse=True)
    links = "".join(f'<li><a href="{d}.html">{d}</a> · <a href="{d}.xlsx">엑셀</a></li>' for d in items)
    (rep_dir / "index.html").write_text(
        f'<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
        f'<title>리포트 목록</title><style>{CSS}</style>{THEME_JS}</head><body><main><h1>일일 리포트 목록</h1>'
        f'<p><a href="../index.html">최신 리포트</a></p><ul>{links}</ul></main></body></html>', encoding="utf-8")
